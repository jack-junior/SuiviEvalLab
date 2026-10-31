# ruff: noqa: E501  (the guide text is prose; long lines keep each sentence searchable)
"""Builds GUIDE-DE-TEST.docx from expected.json (results checked against R survey)."""

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

KIT = Path(sys.argv[1] if len(sys.argv) > 1 else "build/test-kit")
E = json.loads((KIT / "expected.json").read_text())
BLUE = RGBColor(0x1D, 0x4F, 0x91)
MUTED = RGBColor(0x5F, 0x5C, 0x57)


def num(x, d=1):
    s = f"{x:,.{d}f}".replace(",", " ").replace(".", ",")
    return s


def pct(x, d=1):
    return num(x * 100, d) + " %"


doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2)
sec.top_margin = sec.bottom_margin = Cm(1.8)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
for name, size in (("Heading 1", 16), ("Heading 2", 13)):
    h = doc.styles[name]
    h.font.color.rgb = BLUE
    h.font.size = Pt(size)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(rows, header=None, widths=None):
    t = doc.add_table(rows=0, cols=len(header or rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if header:
        cells = t.add_row().cells
        for c, v in zip(cells, header, strict=True):
            c.text = ""
            r = c.paragraphs[0].add_run(v)
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade(c, "1D4F91")
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row, strict=True):
            c.text = str(v)
    if widths:
        for row in t.rows:
            for c, w in zip(row.cells, widths, strict=True):
                c.width = Cm(w)
    doc.add_paragraph()
    return t


def muted(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.color.rgb = MUTED
    r.font.size = Pt(9.5)


n_scen = 0


def scenario(title, goal, files, steps, expected, note=None):
    global n_scen
    n_scen += 1
    doc.add_heading(f"Scénario {n_scen}. {title}", level=2)
    p = doc.add_paragraph()
    p.add_run("Objectif : ").bold = True
    p.add_run(goal)
    if files:
        p = doc.add_paragraph()
        p.add_run("Fichiers : ").bold = True
        p.add_run(files)
    doc.add_paragraph().add_run("Étapes").bold = True
    for i, s in enumerate(steps, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.add_run(f"{i}.  ").bold = True
        p.add_run(s)
    doc.add_paragraph().add_run("Résultat attendu").bold = True
    for e in expected:
        doc.add_paragraph(e, style="List Bullet")
    if note:
        muted(note)
    table(
        [["\n\n", "", ""]],
        header=["Ce que vous avez obtenu", "OK / Problème", "Remarques (ce qui était confus, lent, mal dit)"],
        widths=[6.5, 2.5, 8],
    )


doc.add_heading("SuiviEvalLab 0.3 : kit de test", level=0)
muted("Guide pas à pas pour les testeurs, avec les résultats attendus.")
doc.add_paragraph(
    "Ce kit permet de vérifier l'application comme un utilisateur réel, sur des enquêtes publiques "
    "non modifiées. Chaque scénario donne les étapes, puis le résultat attendu. Les chiffres attendus "
    "ont été calculés une seconde fois sans l'application, directement dans R avec le package survey, "
    "à partir des fichiers bruts : si l'application affiche autre chose, c'est un problème à signaler."
)
doc.add_paragraph(
    "Notez tout ce qui vous gêne, même si le résultat est juste : un bouton difficile à trouver, un mot "
    "peu clair, une attente trop longue. Ce sont des retours aussi utiles que les erreurs de calcul."
)
doc.add_heading("Contenu du kit", level=1)
table(
    [
        [
            "donnees/01_irak_mcna_2019_reference.xlsx",
            "Irak, enquête multisectorielle MCNA VII, 2019 (12 887 ménages, 436 colonnes)",
        ],
        ["donnees/02_irak_mcna_2020_fin_de_projet.xlsx", "Irak, MCNA VIII, 2020 (3 872 ménages, 331 colonnes)"],
        [
            "donnees/03_bangladesh_wash_2018_reference.xlsx",
            "Bangladesh, eau et assainissement dans les camps, avril 2018, avec le questionnaire Kobo",
        ],
        [
            "donnees/04_bangladesh_wash_2018_suivi.xlsx",
            "Bangladesh, suivi de mousson, octobre 2018 (ménages et individus)",
        ],
        [
            "donnees/05_bangladesh_wash_2018_extrait.sav / .dta",
            "Extrait SPSS et Stata (600 ménages, 13 colonnes, avec libellés)",
        ],
        ["modeles/modele_organisation.docx", "Modèle Word d'une organisation fictive, pour le rapport"],
    ],
    header=["Fichier", "Contenu"],
    widths=[7, 10],
)
muted(
    "Les fichiers 01 à 04 sont les fichiers originaux publiés sur Humanitarian Data Exchange, octet pour octet "
    "(empreintes SHA-256 dans SOURCES.txt). Licence CC BY 4.0 : REACH Initiative, Assessment Working Group Iraq, UNICEF."
)
doc.add_heading("Avant de commencer", level=1)
for s in [
    "Installez SuiviEvalLab 0.3.0 (page Releases du dépôt GitHub). Si la 0.2 est installée, lancez simplement le nouvel installeur : inutile de désinstaller.",
    "Copiez ce kit où vous voulez ; gardez les fichiers dans leur dossier « donnees ».",
    "Comptez environ deux heures pour tout le guide. Les scénarios 4 à 11 se suivent ; les autres sont indépendants.",
    "Pour signaler un problème : une capture d'écran, le numéro du scénario et de l'étape, et si possible le fichier journal (chemin indiqué dans l'écran Moteurs).",
]:
    doc.add_paragraph(s, style="List Bullet")

doc.add_heading("Scénarios", level=1)

scenario(
    "Installation et premier lancement",
    "vérifier que l'application s'installe sans droits d'administrateur et que ses moteurs démarrent.",
    "l'installeur SuiviEvalLab-0.3.0-installation.exe.",
    [
        "Lancez l'installeur. Si Windows affiche « Windows a protégé votre ordinateur », cliquez sur « Informations complémentaires » puis « Exécuter quand même ».",
        "Ouvrez SuiviEvalLab. En bas à gauche, attendez « Moteurs prêts ».",
        "Cliquez sur « Moteurs prêts » : l'écran Moteurs s'ouvre.",
    ],
    [
        "Le logo de SuiviEvalLab (carré bleu, courbe blanche) sur le bureau et dans la barre des tâches.",
        "Python 3.12.10 et R 4.4.3 marqués « prêt ».",
        "Une section « Mises à jour » indique la version installée, 0.3.0.",
    ],
)
scenario(
    "Projet d'exemple",
    "voir l'application avec des données déjà prêtes.",
    None,
    ["Accueil, « Ouvrir l'exemple ». Choisissez où enregistrer le fichier.", "Regardez la vue d'ensemble."],
    [
        "Trois indicateurs : école primaire à moins de 2 km « En bonne voie » (85,9 %), école secondaire « En retard », enfants retirés de l'école « Atteint ».",
    ],
)
scenario(
    "Créer un projet et son cadre de résultats",
    "construire un projet vide, comme pour un vrai projet.",
    None,
    [
        "Accueil, « Créer un projet », nommez-le « Test Irak », « Choisir l'emplacement ».",
        "Cadre de résultats : ajoutez un impact « Les enfants ont accès à l'éducation », puis sous lui un effet « L'accès physique aux écoles s'améliore ».",
    ],
    ["Le projet apparaît dans « Mes projets » à l'accueil, avec sa date de modification."],
)
s = E["sheets"]
scenario(
    "Importer l'enquête de référence (fichier original)",
    "importer un vrai fichier volumineux, à plusieurs feuilles.",
    "01_irak_mcna_2019_reference.xlsx",
    [
        "Données, « Importer des données », choisissez le fichier.",
        f"Vérifiez la feuille proposée : « {s['2019']} ».",
        "Collecte : « Référence ». Pondération : la colonne « weights » doit être proposée d'elle-même.",
        "Importez. Notez le temps d'attente.",
    ],
    [
        f"12 887 lignes et {s['cols19']} colonnes.",
        "Plan d'enquête affiché : pondéré par weights.",
    ],
)
q19 = E["quality19"]
scenario(
    "Diagnostic qualité",
    "voir les problèmes trouvés automatiquement dans un fichier réel.",
    "le fichier importé au scénario 4.",
    ["Ouvrez le fichier, onglet « Qualité ».", "Lisez les avertissements. Ne corrigez rien pour l'instant."],
    [
        f"{len(q19)} avertissements « valeurs très éloignées des autres », par exemple tot_child (48 valeurs, jusqu'à 18 enfants) et num_family_member (17 valeurs, jusqu'à 26 personnes).",
        "Une ligne « remarques » repliée en bas (colonnes vides, questions filtrées).",
    ],
    "Ces valeurs sont à vérifier, pas forcément fausses : un ménage de 26 personnes existe. C'est au testeur de dire si le message est clair.",
)
eff = E["effects"]
scenario(
    "Nettoyage : rendre les deux enquêtes comparables",
    "utiliser le journal de nettoyage.",
    "le fichier 2019 importé.",
    [
        "Onglet « Nettoyage », « Ajouter une règle », type « Exclure des lignes » : colonne population_group, valeurs idp_in_camp et host (2020 n'a interrogé ni les camps ni la communauté hôte). Justification : « Comparable à 2020 ».",
        "Ajoutez une règle « Recoder des valeurs » : colonne primary_school_place, valeur do_not_know, vers « valeur manquante ».",
    ],
    [
        f"Première règle : {num(eff['MCNA 2019'][0][1]['rows_removed'], 0)} lignes retirées ; il reste {num(E['clean_rows']['MCNA 2019'], 0)} lignes.",
        f"Seconde règle : {eff['MCNA 2019'][1][1]['cells_changed']} valeurs modifiées.",
        "Le fichier d'origine n'est pas modifié : l'aperçu en « Données brutes » montre toujours les 12 887 lignes.",
    ],
)
scenario(
    "Importer l'enquête de fin de projet",
    "deuxième collecte, avec un plan d'enquête stratifié.",
    "02_irak_mcna_2020_fin_de_projet.xlsx",
    [
        f"Importez le fichier (feuille « {s['2020']} »), collecte « Fin de projet ».",
        "Vérifiez le plan proposé : pondération weights et strates strata.",
        "Ajoutez la même règle « Recoder » do_not_know vers manquant sur primary_school_place.",
        "Ouvrez l'onglet « Qualité ».",
    ],
    [
        f"3 872 lignes, {s['cols20']} colonnes ; la règle modifie {eff['MCNA 2020'][0][1]['cells_changed']} valeurs.",
        "Dans les remarques : la somme des poids (4 440) est proche du nombre de lignes, donc les poids sont normalisés (voir scénario 9).",
    ],
)
r1 = E["ind1"]["rounds"]
c1 = E["ind1"]["comparison"]
by = {g[0]: g[1] for g in r1[0]["by"]}
by2 = {g[0]: g[1] for g in r1[1]["by"]}
scenario(
    "Indicateur en pourcentage, comparaison et taux de réalisation",
    "le cœur de l'application : un indicateur calculé par R avec son intervalle de confiance.",
    "les deux fichiers nettoyés.",
    [
        "Indicateurs, « Ajouter un indicateur » sous l'effet : « Ménages avec une école primaire à moins de 2 km », unité Pourcentage.",
        "Règle : colonne primary_school_place, valeur within_2km. Cible : 90. Désagrégation : population_group.",
        "Ouvrez l'écran Analyse.",
    ],
    [
        f"Référence : {pct(r1[0]['estimate'])} (IC à 95 % : {pct(r1[0]['ci_low'])} à {pct(r1[0]['ci_high'])} ; {num(r1[0]['n'], 0)} observations).",
        f"Fin de projet : {pct(r1[1]['estimate'])} (IC : {pct(r1[1]['ci_low'])} à {pct(r1[1]['ci_high'])} ; {num(r1[1]['n'], 0)} observations).",
        f"Hausse de {num(c1['difference'] * 100)} points (IC : {num(c1['ci_low'] * 100)} à {num(c1['ci_high'] * 100)}), statistiquement significative (p < 0,001).",
        f"Taux de réalisation {pct(E['ind1']['achievement'], 0)}, statut « En bonne voie ».",
        f"Par groupe : idp_out_camp {pct(by['idp_out_camp'])} puis {pct(by2['idp_out_camp'])} ; returnee {pct(by['returnee'])} puis {pct(by2['returnee'])}.",
    ],
)
r2 = E["ind2"]["rounds"]
scenario(
    "Indicateur en effectif : un piège à détecter",
    "vérifier que l'application prévient quand une comparaison n'a pas de sens.",
    "les deux fichiers.",
    [
        "Ajoutez « Ménages ayant retiré un enfant de l'école », unité Effectif, règle child_dropout_school = yes, sens « baisse ».",
        "Ouvrez l'Analyse.",
    ],
    [
        f"Référence : {num(r2[0]['estimate'], 0)} ménages ; fin de projet : {num(r2[1]['estimate'], 0)}.",
        "Un message rouge explique que les poids des deux enquêtes ne sont pas à la même échelle (poids moyen 96,57 contre 1,15) : la baisse n'a aucun sens. Les poids de 2019 extrapolent à la population, ceux de 2020 sont normalisés.",
    ],
    "Ce cas a été trouvé en préparant ce kit sur les fichiers originaux ; l'avertissement a été ajouté à l'application. Pour comparer, la bonne mesure est la part des ménages (pourcentage).",
)
r3 = E["ind3"]["rounds"]
c3 = E["ind3"]["comparison"]
scenario(
    "Indicateur en moyenne",
    "une moyenne pondérée et une différence non significative.",
    None,
    ["Ajoutez « Nombre moyen d'enfants par ménage », unité Moyenne, colonne tot_child.", "Ouvrez l'Analyse."],
    [
        f"Référence {num(r3[0]['estimate'], 2)} (IC {num(r3[0]['ci_low'], 2)} à {num(r3[0]['ci_high'], 2)}), fin {num(r3[1]['estimate'], 2)} (IC {num(r3[1]['ci_low'], 2)} à {num(r3[1]['ci_high'], 2)}).",
        f"Différence non significative (p = {num(c3['p_value'], 3)}).",
    ],
)
te = E["test"]
g = {x[0]: x[1] for x in te["groups"]}
scenario(
    "Assistant de tests",
    "répondre à une question sans connaître le nom du test.",
    "le fichier 2019 nettoyé.",
    [
        "Écran Tests, « Comparer des groupes ». Type : « Oui ou non (pourcentage) ». Variable primary_school_place, valeur within_2km. Groupes : population_group. Fichier : 2019.",
        "Lancez le test, lisez la conclusion, enregistrez-la.",
    ],
    [
        f"returnee {pct(g['returnee'])} contre idp_out_camp {pct(g['idp_out_camp'])}, écart de 1,4 point.",
        f"Test de Rao-Scott (plan d'enquête), p = {num(te['p_value'], 3)} : différence non significative. V de Cramér 0,06 (négligeable).",
        "Le test enregistré apparaîtra dans le rapport Word (« Analyses complémentaires »).",
    ],
)
bl = E["bgd_labels"]
bq = {x[0]: x for x in E["bgd_quality"]}
scenario(
    "Fichier Kobo avec son questionnaire",
    "lire les libellés d'un questionnaire XLSForm et trouver des doublons.",
    "03_bangladesh_wash_2018_reference.xlsx",
    [
        "Créez un nouveau projet « Test Bangladesh ».",
        "Importez le fichier, feuille « Dataset ». Cochez l'utilisation du questionnaire (feuilles Questionnaire et Questionnaire_choices détectées).",
        "Onglet Colonnes puis Aperçu avec « Afficher les libellés ».",
        "Onglet Qualité : corrigez les identifiants en double en un clic.",
    ],
    [
        f"{bl['cols']} colonnes, dont {bl['columns']} avec un libellé lisible du questionnaire.",
        f"Qualité : l'identifiant _id se répète ({bq['duplicate_id'][2]} lignes en double), une réponse hors questionnaire (« Yes » dans la colonne du consentement), 3 tailles de ménage très éloignées (13 et 14).",
        f"Après la correction : {num(E['bgd_after_dedupe'], 0)} ménages, et une règle « Diagnostic qualité : identifiant en double » dans le journal.",
    ],
)
m = E["merge"]["stats"]
scenario(
    "Fusion : chaque contenant d'eau avec son ménage",
    "relier un « repeat » Kobo à son ménage.",
    "03_bangladesh_wash_2018_reference.xlsx, feuille Dataset_containergroup.",
    [
        "Importez la feuille « Dataset_containergroup » du même fichier, nommez-la « Contenants ».",
        "Données, « Fusionner deux fichiers » : fichier principal Contenants, fichier à rattacher le fichier des ménages. Les clés _parent_index et _index sont proposées.",
        "Lisez le contrôle avant de créer, puis créez le fichier fusionné.",
    ],
    [
        f"{num(m['matched'], 0)} contenants sur {num(m['left_rows'], 0)} trouvent leur ménage ; {m['left_unmatched']} sans correspondance ; {m['right_unmatched']} ménages sans contenant.",
        f"{E['merge']['added']} colonnes ajoutées ; celles qui existaient déjà prennent le suffixe _2.",
    ],
    "Si vous n'avez pas corrigé les doublons au scénario précédent, vous obtiendrez 5 438 sur 5 476 : les ménages en double portaient aussi des contenants. C'est normal et c'est justement intéressant à noter.",
)
scenario(
    "Fusion refusée à juste titre",
    "vérifier que l'application refuse une fusion qui dupliquerait des lignes.",
    "04_bangladesh_wash_2018_suivi.xlsx",
    [
        "Importez les feuilles « Household Dataset » et « Individuals Dataset ».",
        "Fusionnez Individus (principal) avec Ménages, clés _parent_index et _index.",
    ],
    [
        "Message : la clé _index se répète dans le second fichier (75 valeurs) ; la fusion n'est pas créée.",
    ],
    "Explication : ce fichier public met bout à bout deux versions du questionnaire Kobo, dont les numéros _index recommencent à 1. Dites-nous si le message vous a permis de comprendre le problème.",
)
scenario(
    "Fichiers SPSS et Stata",
    "importer les formats des logiciels statistiques.",
    "05_bangladesh_wash_2018_extrait.sav puis .dta",
    ["Importez le fichier .sav, onglet Colonnes.", "Faites de même avec le fichier .dta."],
    [
        "600 lignes, 13 colonnes.",
        "La colonne res_age a le libellé « How old are you? » ; dans drnk_wat, le code 2 s'affiche « Tubewells/borehole/handpump ».",
    ],
)
scenario(
    "Rapport, plan de suivi et tableau de bord",
    "produire les documents à partir du projet « Test Irak ».",
    "modeles/modele_organisation.docx",
    [
        "Rouvrez « Test Irak » depuis Mes projets, écran Rapport.",
        "Lisez les vérifications avant le rapport.",
        "Générez le rapport Word, ouvrez-le.",
        "Choisissez le modèle de l'organisation et générez à nouveau.",
        "Exportez le plan de suivi (Excel), puis créez le tableau de bord (HTML) et ouvrez-le dans un navigateur.",
    ],
    [
        "Vérifications : deux indicateurs sans cible, et l'avertissement sur les poids de l'indicateur en effectif.",
        "Rapport : synthèse, résultats par effet avec graphiques, test enregistré, annexes données et méthodes.",
        "Avec le modèle : page de garde et styles de l'organisation, le contenu à l'emplacement {{contenu}}.",
        "Plan de suivi : une ligne par indicateur, mode de calcul en phrases.",
        "Tableau de bord : un seul fichier, filtres par statut, aucune ligne de données.",
    ],
)
scenario(
    "Taille d'échantillon",
    "préparer une enquête.",
    None,
    [
        "Accueil, « Calculer une taille d'échantillon ».",
        "Estimer une part : 50 %, marge 5 points, effet de plan 1,5, taux de réponse 90 %.",
        "Détecter un changement de part : de 50 % à 60 %, mêmes réglages, puissance 80 %.",
    ],
    [f"{E['sample1']} enquêtés au total.", f"{E['sample2']} enquêtés à chaque collecte."],
    "Vérification à la main : 1,96² × 0,5 × 0,5 / 0,05² = 384,1 ; × 1,5 = 576,2 ; / 0,9 = 640,2, arrondi à 641.",
)
scenario(
    "Copies automatiques",
    "revenir en arrière après une erreur.",
    None,
    [
        "Fermez « Test Irak » et rouvrez-le (une copie est faite à l'ouverture).",
        "Supprimez l'effet du cadre de résultats (avec ses indicateurs). Fermez le projet.",
        "Accueil, icône horloge du projet : choisissez la copie la plus récente, « Restaurer ».",
    ],
    [
        "Le projet revient avec son effet et ses indicateurs. Une nouvelle copie de l'état d'avant la restauration apparaît dans la liste."
    ],
)
scenario(
    "Langue et mises à jour",
    "dernières vérifications.",
    None,
    [
        "En bas à gauche, passez en EN puis revenez en FR.",
        "Écran Moteurs, section Mises à jour : décochez puis recochez la recherche automatique, cliquez « Vérifier maintenant ».",
    ],
    [
        "Toute l'interface change de langue, y compris les dates et les nombres.",
        "Message « L'application est à jour » tant qu'aucune version plus récente n'est publiée. À la prochaine version, un bandeau proposera de redémarrer pour l'installer.",
    ],
)

doc.add_heading("Comment les résultats attendus ont été obtenus", level=1)
doc.add_paragraph(
    "Chaque scénario a d'abord été joué automatiquement avec le moteur de l'application, exactement comme décrit. "
    "Les chiffres des scénarios 8 à 11 ont ensuite été recalculés indépendamment, dans R, à partir des fichiers "
    "bruts : svydesign avec les poids (et les strates en 2020), svyciprop méthode logit pour les parts, svytotal, "
    "svymean, différence avec erreur type combinée et test z, svychisq (Rao-Scott) pour le test entre groupes. "
    "Les deux calculs concordent jusqu'à la dernière décimale affichée."
)
doc.add_heading("Signaler un problème", level=1)
for s in [
    "Sur GitHub : https://github.com/jack-junior/SuiviEvalLab/issues, ou directement au chef de produit.",
    "Indiquez le scénario et l'étape, ce que vous attendiez, ce que vous avez obtenu, avec une capture d'écran.",
    "Joignez si possible le journal : son dossier est indiqué dans l'écran Moteurs.",
]:
    doc.add_paragraph(s, style="List Bullet")

doc.save(KIT / "GUIDE-DE-TEST.docx")
print("ok", n_scen)
