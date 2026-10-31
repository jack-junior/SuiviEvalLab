"""Build tests/fixtures/modele_organisation.docx: an organisation template with its own fonts,
colours, header, footer and cover fields, used by the report tests.

    python tools/datasets/make_report_template.py engine/python/tests/fixtures/modele_organisation.docx
"""

import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

doc = Document()
st = doc.styles
st["Normal"].font.name = "Arial"
st["Normal"].font.size = Pt(10)
for n in ("Heading 1", "Heading 2", "Heading 3"):
    st[n].font.name = "Arial"
    st[n].font.color.rgb = RGBColor(0x0B, 0x6E, 0x4F)
sec = doc.sections[0]
sec.header.paragraphs[0].text = "ONG Exemple, Programme Éducation  |  {{projet}}"
sec.footer.paragraphs[0].text = "Document interne, {{date}}"
t = doc.add_paragraph()
r = t.add_run("{{projet}}")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = RGBColor(0x0B, 0x6E, 0x4F)
doc.add_paragraph("{{titre_rapport}}")
doc.add_paragraph("Préparé par l'équipe suivi-évaluation, {{date}}").alignment = WD_ALIGN_PARAGRAPH.LEFT
doc.add_page_break()
doc.add_heading("Résumé exécutif", level=1)
doc.add_paragraph("Texte rédigé par l'équipe avant les résultats.")
doc.add_paragraph("{{contenu}}")
doc.add_heading("Recommandations", level=1)
doc.add_paragraph("Texte rédigé par l'équipe après les résultats.")
doc.save(sys.argv[1])
