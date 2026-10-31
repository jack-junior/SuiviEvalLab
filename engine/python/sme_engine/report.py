"""Word report generated from the project: summary, results by outcome, appendices.

The text follows the writing conventions: magnitude first, then significance,
numbers in the report language, no decorative punctuation.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import __version__, xlsform

INK = RGBColor(0x1F, 0x1E, 0x1C)
ACTION = RGBColor(0x1D, 0x4F, 0x91)
MUTED = RGBColor(0x5F, 0x5C, 0x57)

T = {
    "fr": {
        "subtitle": "Rapport de suivi des indicateurs",
        "generated": "Généré le {date} avec SuiviEvalLab {version}",
        "summary": "Synthèse",
        "summary_text": "Indicateurs suivis : {n}. Atteints : {achieved}. En bonne voie : {on_track}. En retard : {late}. Non évalués : {unknown}.",
        "col_indicator": "Indicateur",
        "col_baseline": "Référence",
        "col_current": "Valeur actuelle",
        "col_target": "Cible",
        "col_rate": "Réalisation",
        "col_status": "Statut",
        "status": {"achieved": "Atteint", "on_track": "En bonne voie", "late": "En retard", "unknown": "Non évalué"},
        "results": "Résultats par résultat attendu",
        "levels": {"impact": "Impact", "outcome": "Effet", "output": "Extrant", "activity": "Activité"},
        "definition": "Définition : ",
        "source": "Source : ",
        "no_data": "Aucune donnée disponible pour cet indicateur.",
        "rounds": {
            "baseline": "Référence",
            "midline": "Mi-parcours",
            "endline": "Fin de projet",
            "monitoring": "Suivi",
        },
        "round_phrase": {
            "baseline": "lors de l'enquête de référence",
            "midline": "à mi-parcours",
            "endline": "lors de l'enquête de fin de projet",
            "monitoring": "lors du suivi",
        },
        "single": "{name} est de {value} {when} (IC à 95 % : {low} à {high} ; {n} observations).",
        "change": "{name} est passé de {a} {when_a} à {b} {when_b}, soit {direction} de {diff} (IC à 95 % : {low} à {high}).",
        "change_name_share": "La part des {name}",
        "up": "une hausse",
        "down": "une baisse",
        "points": "points",
        "sig": "Cette différence est statistiquement significative ({p}).",
        "nonsig": "Cette différence n'est pas statistiquement significative ({p}) : elle peut s'expliquer par les variations d'échantillonnage.",
        "target_rate": "Par rapport à la cible de {target}, le taux de réalisation est de {rate}.",
        "target_reached": "La cible de {target} est atteinte.",
        "by": "Désagrégation par {column}",
        "col_group": "Groupe",
        "chart_target": "Cible",
        "chart_ylab_percent": "% des observations",
        "framework": "Cadre de résultats",
        "col_code": "Code",
        "col_level": "Niveau",
        "col_title": "Intitulé",
        "col_indicators": "Indicateurs",
        "annex_data": "Annexe A. Données et nettoyage",
        "ds_line": "Fichier {file} ({sheet}), importé le {date} : {rows} lignes, {clean} après nettoyage.",
        "design": "Plan d'enquête : {text}.",
        "design_none": "échantillon aléatoire simple, sans pondération",
        "design_weights": "pondération par « {w} »",
        "design_strata": "strates « {s} »",
        "design_cluster": "grappes « {c} »",
        "rules": "Journal de nettoyage",
        "col_rule": "Règle",
        "col_effect": "Effet",
        "col_note": "Justification",
        "rule_text": {
            "duplicates": "Doublons sur {columns} : garder la première ligne",
            "recode": "Dans « {column} », remplacer {from_} par {to}",
            "range": "Dans « {column} », valeurs hors de [{min}, {max}] mises en manquant",
            "exclude": "Exclure les lignes où « {column} » vaut {values}",
        },
        "missing": "valeur manquante",
        "effect_rows": "{n} lignes retirées",
        "effect_cells": "{n} valeurs modifiées",
        "effect_none": "aucun changement",
        "disabled": "désactivée",
        "annex_methods": "Annexe B. Méthodes statistiques",
        "methods": {
            "wilson": "Proportion avec intervalle de confiance de Wilson (échantillon aléatoire simple).",
            "survey_logit": "Proportion pondérée selon le plan d'enquête, intervalle de confiance par la méthode logit (package R survey, linéarisation de Taylor).",
            "t_mean": "Moyenne avec intervalle de confiance de Student (échantillon aléatoire simple).",
            "survey_mean": "Moyenne pondérée selon le plan d'enquête, intervalle de confiance par linéarisation de Taylor (package R survey).",
        },
        "tests": {
            "z_survey_prop": "Différence entre deux enquêtes indépendantes, chacune analysée avec son plan d'enquête ; erreur type combinée et test z.",
            "z_survey_mean": "Différence de moyennes entre deux enquêtes indépendantes, chacune analysée avec son plan d'enquête ; erreur type combinée et test z.",
            "chisq_prop": "Test du khi-deux de comparaison de deux proportions (sans correction de continuité).",
            "welch": "Test t de Welch de comparaison de deux moyennes.",
        },
        "methods_intro": "Tous les calculs sont faits par R. Les résultats sont accompagnés de leur intervalle de confiance à 95 % ; les comparaisons donnent la différence, son intervalle, la valeur p et une taille d'effet ({effect}).",
        "effect_names": "h de Cohen pour les proportions, d de Cohen pour les moyennes",
        "achievement_def": "Taux de réalisation = (valeur actuelle − référence) / (cible − référence). Seuils : 90 % et plus atteint, 60 % à 90 % en bonne voie, moins de 60 % en retard.",
        "sources": "Sources des données",
        "p_small": "p < 0,001",
        "p_eq": "p = {p}",
    },
    "en": {
        "subtitle": "Indicator monitoring report",
        "generated": "Generated on {date} with SuiviEvalLab {version}",
        "summary": "Summary",
        "summary_text": "Indicators tracked: {n}. Achieved: {achieved}. On track: {on_track}. Late: {late}. Not assessed: {unknown}.",
        "col_indicator": "Indicator",
        "col_baseline": "Baseline",
        "col_current": "Current value",
        "col_target": "Target",
        "col_rate": "Achievement",
        "col_status": "Status",
        "status": {"achieved": "Achieved", "on_track": "On track", "late": "Late", "unknown": "Not assessed"},
        "results": "Results by expected result",
        "levels": {"impact": "Impact", "outcome": "Outcome", "output": "Output", "activity": "Activity"},
        "definition": "Definition: ",
        "source": "Source: ",
        "no_data": "No data available for this indicator.",
        "rounds": {"baseline": "Baseline", "midline": "Midline", "endline": "Endline", "monitoring": "Monitoring"},
        "round_phrase": {
            "baseline": "at baseline",
            "midline": "at midline",
            "endline": "at endline",
            "monitoring": "at monitoring",
        },
        "single": "{name} is {value} {when} (95% CI: {low} to {high}; {n} observations).",
        "change": "{name} went from {a} {when_a} to {b} {when_b}, {direction} of {diff} (95% CI: {low} to {high}).",
        "change_name_share": "The share of {name}",
        "up": "an increase",
        "down": "a decrease",
        "points": "points",
        "sig": "This difference is statistically significant ({p}).",
        "nonsig": "This difference is not statistically significant ({p}): it may be due to sampling variation.",
        "target_rate": "Against the target of {target}, the achievement rate is {rate}.",
        "target_reached": "The target of {target} is reached.",
        "by": "Breakdown by {column}",
        "col_group": "Group",
        "chart_target": "Target",
        "chart_ylab_percent": "% of observations",
        "framework": "Results framework",
        "col_code": "Code",
        "col_level": "Level",
        "col_title": "Title",
        "col_indicators": "Indicators",
        "annex_data": "Appendix A. Data and cleaning",
        "ds_line": "File {file} ({sheet}), imported on {date}: {rows} rows, {clean} after cleaning.",
        "design": "Survey design: {text}.",
        "design_none": "simple random sample, unweighted",
        "design_weights": 'weights "{w}"',
        "design_strata": 'strata "{s}"',
        "design_cluster": 'clusters "{c}"',
        "rules": "Cleaning journal",
        "col_rule": "Rule",
        "col_effect": "Effect",
        "col_note": "Reason",
        "rule_text": {
            "duplicates": "Duplicates on {columns}: keep the first row",
            "recode": 'In "{column}", replace {from_} with {to}',
            "range": 'In "{column}", values outside [{min}, {max}] set to missing',
            "exclude": 'Exclude rows where "{column}" is {values}',
        },
        "missing": "missing value",
        "effect_rows": "{n} rows removed",
        "effect_cells": "{n} values changed",
        "effect_none": "no change",
        "disabled": "disabled",
        "annex_methods": "Appendix B. Statistical methods",
        "methods": {
            "wilson": "Proportion with Wilson confidence interval (simple random sample).",
            "survey_logit": "Proportion weighted by the survey design, logit confidence interval (R survey package, Taylor linearisation).",
            "t_mean": "Mean with Student confidence interval (simple random sample).",
            "survey_mean": "Mean weighted by the survey design, confidence interval by Taylor linearisation (R survey package).",
        },
        "tests": {
            "z_survey_prop": "Difference between two independent surveys, each analysed with its own design; combined standard error and z test.",
            "z_survey_mean": "Difference of means between two independent surveys, each analysed with its own design; combined standard error and z test.",
            "chisq_prop": "Chi-squared test comparing two proportions (no continuity correction).",
            "welch": "Welch t test comparing two means.",
        },
        "methods_intro": "All calculations are done in R. Results come with their 95% confidence interval; comparisons give the difference, its interval, the p-value and an effect size ({effect}).",
        "effect_names": "Cohen's h for proportions, Cohen's d for means",
        "achievement_def": "Achievement rate = (current value − baseline) / (target − baseline). Thresholds: 90% and above achieved, 60% to 90% on track, below 60% late.",
        "sources": "Data sources",
        "p_small": "p < 0.001",
        "p_eq": "p = {p}",
    },
}


class Fmt:
    def __init__(self, lang: str):
        self.lang = lang
        self.none = "nd" if lang == "fr" else "n/a"  # no dash for empty values (writing conventions)

    def num(self, x: float | None, digits: int = 1) -> str:
        if x is None:
            return self.none
        s = f"{x:,.{digits}f}"
        if self.lang == "fr":
            s = s.replace(",", " ").replace(".", ",")
        return s

    def pct(self, share: float | None, digits: int = 1) -> str:
        if share is None:
            return self.none
        return self.num(share * 100, digits) + (" %" if self.lang == "fr" else "%")

    def value(self, x: float | None, unit: str) -> str:
        return self.pct(x) if unit == "percent" else self.num(x, 2)

    def points(self, diff: float, unit: str, t: dict[str, Any]) -> str:
        return f"{self.num(abs(diff) * 100)} {t['points']}" if unit == "percent" else self.num(abs(diff), 2)

    def p(self, p: float, t: dict[str, Any]) -> str:
        return t["p_small"] if p < 0.001 else t["p_eq"].format(p=self.num(p, 3))


def _shade(cell: Any, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _table(doc: Any, header: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        _shade(cell, "EEEDE9")
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(v)
            run.font.size = Pt(9)
            if i > 0:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    if widths:
        table.autofit = False
        grid = table._tbl.tblGrid
        for i, col in enumerate(grid.findall(qn("w:gridCol"))):
            col.set(qn("w:w"), str(int(Cm(widths[i]).twips)))
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def _setup(doc: Any) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        st = styles[name]
        st.font.name = "Cambria"
        st.font.size = Pt(size)
        st.font.color.rgb = INK if name != "Heading 1" else ACTION
        st.font.bold = name != "Title"
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2)


def _design_text(design: dict[str, Any], t: dict[str, Any]) -> str:
    parts = []
    if design.get("weights"):
        parts.append(t["design_weights"].format(w=design["weights"]))
    if design.get("strata"):
        parts.append(t["design_strata"].format(s=" × ".join(design["strata"])))
    if design.get("cluster"):
        parts.append(t["design_cluster"].format(c=design["cluster"]))
    return ", ".join(parts) if parts else t["design_none"]


def _rule_text(rule: dict[str, Any], t: dict[str, Any], f: Fmt) -> str:
    q = (lambda v: f"« {v} »") if f.lang == "fr" else (lambda v: f'"{v}"')
    kind = rule["kind"]
    tpl = t["rule_text"][kind]
    if kind == "duplicates":
        return tpl.format(columns=", ".join(q(c) for c in rule["columns"]))
    if kind == "recode":
        to = t["missing"] if rule["to"] is None else q(rule["to"])
        return tpl.format(column=rule["column"], from_=", ".join(q(v) for v in rule["from"]), to=to)
    if kind == "range":
        return tpl.format(
            column=rule["column"],
            min=f.num(rule["min"], 2) if rule["min"] is not None else "−∞",
            max=f.num(rule["max"], 2) if rule["max"] is not None else "+∞",
        )
    return tpl.format(
        column=rule["column"], values=", ".join(t["missing"] if v is None else q(v) for v in rule["values"])
    )


def _effect_text(rule: dict[str, Any], t: dict[str, Any]) -> str:
    e = rule.get("effect") or {}
    if not rule.get("enabled", True):
        return t["disabled"]
    parts = []
    if e.get("rows_removed"):
        parts.append(t["effect_rows"].format(n=e["rows_removed"]))
    if e.get("cells_changed"):
        parts.append(t["effect_cells"].format(n=e["cells_changed"]))
    return ", ".join(parts) or t["effect_none"]


def _lower_first(text: str) -> str:
    return text[:1].lower() + text[1:] if text and not text[:2].isupper() else text


def _narrative(ind: dict[str, Any], an: dict[str, Any], t: dict[str, Any], f: Fmt) -> list[str]:
    unit = ind["unit"]
    usable = [r for r in an["rounds"] if "error_key" not in r["result"]]
    name = t["change_name_share"].format(name=_lower_first(ind["name"])) if unit == "percent" else ind["name"]
    out = []
    comp = an.get("comparison")
    if comp and "error_key" not in comp:
        a = next(r for r in usable if r["dataset_id"] == comp["a_id"])
        b = next(r for r in usable if r["dataset_id"] == comp["b_id"])
        out.append(
            t["change"].format(
                name=name,
                a=f.value(comp["a"]["estimate"], unit),
                when_a=t["round_phrase"][a["round"]],
                b=f.value(comp["b"]["estimate"], unit),
                when_b=t["round_phrase"][b["round"]],
                direction=t["up"] if comp["difference"] >= 0 else t["down"],
                diff=f.points(comp["difference"], unit, t),
                low=(f.num(comp["ci_low"] * 100) if unit == "percent" else f.num(comp["ci_low"], 2)),
                high=(f.num(comp["ci_high"] * 100) if unit == "percent" else f.num(comp["ci_high"], 2)),
            )
        )
        out.append((t["sig"] if comp["p_value"] < 0.05 else t["nonsig"]).format(p=f.p(comp["p_value"], t)))
    elif usable:
        r = usable[-1]
        res = r["result"]
        out.append(
            t["single"].format(
                name=name,
                value=f.value(res["estimate"], unit),
                when=t["round_phrase"][r["round"]],
                low=f.value(res.get("ci_low"), unit),
                high=f.value(res.get("ci_high"), unit),
                n=f.num(res["n"], 0),
            )
        )
    if an.get("target") is not None and an.get("current") is not None:
        if an["status"] == "achieved" and (an.get("achievement") is None or an["achievement"] >= 1):
            out.append(t["target_reached"].format(target=f.value(an["target"], unit)))
        elif an.get("achievement") is not None:
            out.append(t["target_rate"].format(target=f.value(an["target"], unit), rate=f.pct(an["achievement"], 0)))
    return out


def generate(svc: Any, path: Path, lang: str) -> dict[str, Any]:
    t = T[lang]
    f = Fmt(lang)
    snap = svc.snapshot()
    project = snap["project"]
    nodes = {n["id"]: n for n in snap["framework"]}
    analyses = {ind["id"]: svc.indicator_analysis(ind["id"]) for ind in snap["indicators"]}

    def readable(column: str, code: str | None = None) -> str:
        """Readable label of a column or of one of its codes, from the first dataset that has one."""
        for ds in snap["datasets"]:
            labels = ds.get("labels") or {}
            found = (
                xlsform.pick(labels.get("values", {}).get(column, {}).get(code), lang)
                if code is not None
                else xlsform.pick(labels.get("columns", {}).get(column), lang)
            )
            if found:
                return found
        return code if code is not None else column

    doc = Document()
    _setup(doc)
    doc.add_paragraph(project["name"], style="Title")
    sub = doc.add_paragraph(t["subtitle"])
    sub.runs[0].font.size = Pt(13)
    sub.runs[0].font.color.rgb = MUTED
    gen = doc.add_paragraph(t["generated"].format(date=f_date(lang), version=__version__))
    gen.runs[0].font.size = Pt(9)
    gen.runs[0].font.color.rgb = MUTED

    # 1. Summary
    doc.add_heading(f"1. {t['summary']}", level=1)
    counts = {"achieved": 0, "on_track": 0, "late": 0, "unknown": 0}
    rows = []
    for ind in snap["indicators"]:
        an = analyses[ind["id"]]
        counts[an["status"]] += 1
        node = nodes.get(ind["node_id"], {})
        rows.append(
            [
                f"{node.get('code', '')} {ind['name']}".strip(),
                f.value(an["baseline"], ind["unit"]),
                f.value(an["current"], ind["unit"]),
                f.value(an["target"], ind["unit"]),
                f.pct(an["achievement"], 0),
                t["status"][an["status"]],
            ]
        )
    doc.add_paragraph(t["summary_text"].format(n=len(rows), **counts))
    _table(
        doc,
        [t["col_indicator"], t["col_baseline"], t["col_current"], t["col_target"], t["col_rate"], t["col_status"]],
        rows,
        [6.4, 2, 2.2, 1.8, 2, 2.2],
    )

    # 2. Results
    doc.add_heading(f"2. {t['results']}", level=1)
    workdir = Path(svc.bridge.workdir)
    for node in snap["framework"]:
        inds = [i for i in snap["indicators"] if i["node_id"] == node["id"]]
        if not inds:
            continue
        label = t["levels"][node["level"]]
        doc.add_heading(
            f"{label} {node['code']} : {node['title']}"
            if lang == "fr" and node["code"]
            else f"{label} {node['code']}: {node['title']}"
            if node["code"]
            else f"{label}: {node['title']}",
            level=2,
        )
        for ind in inds:
            an = analyses[ind["id"]]
            doc.add_heading(ind["name"], level=3)
            if ind.get("definition"):
                p = doc.add_paragraph()
                p.add_run(t["definition"]).bold = True
                p.add_run(ind["definition"])
            usable = [r for r in an["rounds"] if "error_key" not in r["result"]]
            if not usable:
                doc.add_paragraph(t["no_data"])
                continue
            for sentence in [" ".join(_narrative(ind, an, t, f))]:
                doc.add_paragraph(sentence)
            # Chart
            png = workdir / f"chart_{ind['id']}.png"
            svc.bridge.call(
                "chart",
                {
                    "out": str(png),
                    "labels": [t["rounds"][r["round"]] for r in usable],
                    "values": [r["result"]["estimate"] for r in usable],
                    "lows": [r["result"].get("ci_low") for r in usable],
                    "highs": [r["result"].get("ci_high") for r in usable],
                    "target": an.get("target"),
                    "percent": ind["unit"] == "percent",
                    "title": "",
                    "ylab": t["chart_ylab_percent"] if ind["unit"] == "percent" else (ind.get("unit_label") or ""),
                    "target_label": t["chart_target"],
                    "decimal": "," if lang == "fr" else ".",
                },
            )
            doc.add_picture(str(png), width=Cm(15))
            # Breakdown
            if ind.get("disaggregation") and any(r["result"].get("by") for r in usable):
                by_col = ind["disaggregation"]
                doc.add_paragraph(t["by"].format(column=readable(by_col))).runs[0].bold = True
                groups = sorted({g["group"] for r in usable for g in r["result"].get("by", [])})
                body = []
                for g in groups:
                    line = [readable(by_col, str(g))]
                    for r in usable:
                        hit = next((x for x in r["result"].get("by", []) if x["group"] == g), None)
                        line.append(
                            f"{f.value(hit.get('estimate'), ind['unit'])} (n = {f.num(hit['n'], 0)})"
                            if hit and hit.get("estimate") is not None
                            else f.none
                        )
                    body.append(line)
                _table(doc, [t["col_group"], *[t["rounds"][r["round"]] for r in usable]], body)
            if ind.get("source"):
                p = doc.add_paragraph()
                r1 = p.add_run(t["source"])
                r1.font.size = Pt(9)
                r1.font.color.rgb = MUTED
                r2 = p.add_run(ind["source"])
                r2.font.size = Pt(9)
                r2.font.color.rgb = MUTED

    # 3. Framework
    doc.add_heading(f"3. {t['framework']}", level=1)
    fw_rows = []
    for node in snap["framework"]:
        names = "; ".join(i["name"] for i in snap["indicators"] if i["node_id"] == node["id"])
        fw_rows.append([node["code"] or "", t["levels"][node["level"]], node["title"], names])
    _table(doc, [t["col_code"], t["col_level"], t["col_title"], t["col_indicators"]], fw_rows, [1.4, 2, 7.8, 5.4])

    # Appendix A
    doc.add_heading(t["annex_data"], level=1)
    for ds in snap["datasets"]:
        doc.add_heading(f"{ds['name']} ({t['rounds'][ds['round']]})", level=2)
        doc.add_paragraph(
            t["ds_line"].format(
                file=ds["source_name"],
                sheet=ds.get("sheet") or "CSV",
                date=f_date(lang, ds["imported_at"][:10]),
                rows=f.num(ds["rows"], 0),
                clean=f.num(ds["clean_rows"], 0),
            )
        )
        doc.add_paragraph(t["design"].format(text=_design_text(ds["design"], t)))
        if ds["rules"]:
            doc.add_paragraph(t["rules"]).runs[0].bold = True
            _table(
                doc,
                ["#", t["col_rule"], t["col_effect"], t["col_note"]],
                [
                    [str(i), _rule_text(r, t, f), _effect_text(r, t), r.get("note") or ""]
                    for i, r in enumerate(ds["rules"], start=1)
                ],
                [0.8, 7, 3.6, 5.2],
            )

    # Appendix B
    doc.add_heading(t["annex_methods"], level=1)
    doc.add_paragraph(t["methods_intro"].format(effect=t["effect_names"]))
    used_methods, used_tests = set(), set()
    for an in analyses.values():
        for r in an["rounds"]:
            if "method_key" in r["result"]:
                used_methods.add(r["result"]["method_key"])
        if an.get("comparison") and "test_key" in an["comparison"]:
            used_tests.add(an["comparison"]["test_key"])
    for key in sorted(used_methods):
        doc.add_paragraph(t["methods"][key], style="List Bullet")
    for key in sorted(used_tests):
        doc.add_paragraph(t["tests"][key], style="List Bullet")
    doc.add_paragraph(t["achievement_def"])

    attributions = [ds["attribution"] for ds in snap["datasets"] if ds.get("attribution")]
    if attributions:
        doc.add_heading(t["sources"], level=1)
        for a in dict.fromkeys(attributions):
            doc.add_paragraph(a, style="List Bullet")

    if path.suffix.lower() != ".docx":
        path = path.with_suffix(".docx")
    doc.save(path)
    return {"path": str(path), "indicators": len(rows)}


def f_date(lang: str, iso: str | None = None) -> str:
    d = date.fromisoformat(iso) if iso else date.today()
    if lang == "fr":
        months = [
            "janvier",
            "février",
            "mars",
            "avril",
            "mai",
            "juin",
            "juillet",
            "août",
            "septembre",
            "octobre",
            "novembre",
            "décembre",
        ]
        return f"{d.day} {months[d.month - 1]} {d.year}"
    months = [
        "January",
        "February",
        "March",
        "April",
        "May",
        "June",
        "July",
        "August",
        "September",
        "October",
        "November",
        "December",
    ]
    return f"{d.day} {months[d.month - 1]} {d.year}"
