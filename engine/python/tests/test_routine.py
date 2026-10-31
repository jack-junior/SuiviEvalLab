"""Framework models and routine monitoring indicators."""

from __future__ import annotations

import io
import json
import sqlite3

import pytest
from docx import Document

from sme_engine import frameworks, routine
from sme_engine.methods import register
from sme_engine.r_bridge import RBridge, find_rscript
from sme_engine.rpc import Server

needs_r = pytest.mark.skipif(find_rscript() is None, reason="R not installed")


@pytest.fixture()
def call():
    bridge = RBridge()
    server = Server(stdin=io.StringIO(), stdout=io.StringIO())
    register(server, bridge)

    def _call(method, params=None):
        resp = server.handle_line(json.dumps({"jsonrpc": "2.0", "id": 1, "method": method, "params": params or {}}))
        if "error" in resp:
            raise RuntimeError(resp["error"]["data"]["key"] + ": " + resp["error"]["message"])
        return resp["result"]

    yield _call
    bridge.stop()


# ----- framework models ----------------------------------------------------------


def test_models_and_custom_validation():
    assert [lv["fr"] for lv in frameworks.MODELS["eu"]][:2] == ["Objectif global", "Objectif spécifique"]
    custom = frameworks.validate(
        {"id": "custom", "levels": [{"fr": "But", "en": "Goal"}, {"fr": "Résultat", "indicators": True}]}
    )
    assert custom["levels"][1]["en"] == "Résultat"
    for bad in (
        {"id": "x"},
        {"id": "custom", "levels": [{"fr": "Seul"}]},
        {"id": "custom", "levels": [{"fr": ""}] * 2},
    ):
        with pytest.raises(Exception, match="framework|level|name"):
            frameworks.validate(bad)


def test_usaid_and_world_bank_projects(call, tmp_path):
    snap = call("project.create", {"path": str(tmp_path / "u.sme"), "name": "U", "model": {"id": "usaid"}})
    assert snap["project"]["framework_model"]["id"] == "usaid"
    goal = call("framework.add", {"title": "Goal", "assumptions": "Stabilité politique"})["created"]
    do = call("framework.add", {"parent_id": goal, "title": "DO 1"})["created"]
    ir = call("framework.add", {"parent_id": do, "title": "IR 1.1"})["created"]
    sub = call("framework.add", {"parent_id": ir, "title": "Sub-IR 1.1.1"})["created"]
    with pytest.raises(RuntimeError, match="node_no_children"):
        call("framework.add", {"parent_id": sub, "title": "Too deep"})
    # USAID: indicators on every level, including the fourth.
    call("indicator.save", {"node_id": sub, "indicator": {"name": "Ind", "unit": "count", "source_type": "routine"}})
    fw = call("project.get")["framework"]
    assert fw[0]["assumptions"] == "Stabilité politique"

    # Back to the standard model: the fourth level then carries no indicator, so it is refused.
    with pytest.raises(RuntimeError, match="framework_model_indicators"):
        call("framework.set_model", {"model": {"id": "standard"}})
    # The World Bank model has two levels only: refused while four are used.
    with pytest.raises(RuntimeError, match="framework_model_too_shallow"):
        call("framework.set_model", {"model": {"id": "worldbank"}})
    snap = call(
        "framework.set_model",
        {"model": {"id": "custom", "levels": [{"fr": "A"}, {"fr": "B"}, {"fr": "C"}, {"fr": "D"}]}},
    )
    assert snap["project"]["framework_model"]["levels"][3]["fr"] == "D"


def test_migration_from_version_2(tmp_path, call):
    path = tmp_path / "old.sme"
    call("project.create", {"path": str(path), "name": "Old"})
    call("framework.add", {"title": "Impact"})
    call("project.close")
    conn = sqlite3.connect(path)
    conn.executescript(
        "ALTER TABLE framework_node DROP COLUMN assumptions; DROP TABLE routine_value; PRAGMA user_version = 2;"
    )
    conn.close()
    snap = call("project.open", {"path": str(path)})
    assert snap["framework"][0]["assumptions"] == "" and snap["project"]["framework_model"]["id"] == "standard"
    assert sqlite3.connect(path).execute("PRAGMA user_version").fetchone()[0] == 3


# ----- routine values ------------------------------------------------------------


def test_cumulative_count_with_plan_and_categories():
    ind = {"unit": "count", "aggregation": "sum", "target": 1500, "categories": ["Femmes", "Hommes"]}
    rows = routine.validate_values(
        [
            {"period": "2026-T2", "end": "2026-06-30", "by": {"Femmes": "300", "Hommes": "250"}, "planned": 400},
            {"period": "2026-T1", "end": "2026-03-31", "achieved": "320", "planned": "300", "by": {"Femmes": 180}},
            {"period": "2026-T3", "end": "2026-09-30", "planned": 400},
        ],
        ind["categories"],
        percent=False,
    )
    assert [r["period"] for r in rows] == ["2026-T1", "2026-T2", "2026-T3"]  # sorted by end date
    assert rows[1]["achieved"] == 550  # sum of the categories
    an = routine.analyse(ind, rows)
    assert an["current"] == 870 and an["planned_to_date"] == 700 and an["current_period"] == "2026-T2"
    assert an["achievement"] == pytest.approx(870 / 1500)
    assert an["status"] == "on_track"  # ahead of the plan to date
    assert an["by_total"] == {"Femmes": 480, "Hommes": 250}
    assert an["periods"][2]["to_date"] is None and an["periods"][2]["planned_to_date"] == 1100


def test_rate_takes_the_latest_value_and_late_status():
    ind = {"unit": "percent", "aggregation": "last", "target": 0.8, "baseline_manual": 0.5}
    rows = routine.validate_values(
        [{"period": "M1", "achieved": "55", "planned": 60}, {"period": "M2", "achieved": 0.58, "planned": "70"}],
        [],
        percent=True,
    )
    an = routine.analyse(ind, rows)
    assert an["current"] == pytest.approx(0.58) and an["planned_to_date"] == pytest.approx(0.7)
    assert an["status"] == "late"  # 0.58 / 0.70 = 83 % of the plan
    assert an["achievement"] == pytest.approx((0.58 - 0.5) / (0.8 - 0.5))


def test_decrease_and_errors():
    ind = {"unit": "count", "aggregation": "last", "target": 10, "direction": "decrease"}
    an = routine.analyse(ind, routine.validate_values([{"period": "P1", "achieved": 8}], [], False))
    assert an["status"] == "achieved"
    with pytest.raises(Exception, match="twice"):
        routine.validate_values([{"period": "A", "achieved": 1}, {"period": "A", "achieved": 2}], [], False)
    with pytest.raises(Exception, match="number"):
        routine.validate_values([{"period": "A", "achieved": "beaucoup"}], [], False)
    with pytest.raises(Exception, match="YYYY-MM-DD"):
        routine.validate_values([{"period": "A", "end": "31/03/2026"}], [], False)


def test_paste_from_excel():
    text = (
        "Période\tFin\tRéalisé\tPrévu\tFemmes\tHommes\n"
        "2026-T1\t2026-03-31\t\t300\t180\t140\n"
        "2026-T2\t2026-06-30\t1 250,5\t400\t\t\n"
    )
    rows = routine.validate_values(routine.parse_paste(text, ["Femmes", "Hommes"]), ["Femmes", "Hommes"], False)
    assert rows[0]["achieved"] == 320 and rows[1]["achieved"] == 1250.5 and rows[1]["planned"] == 400


@needs_r
def test_routine_indicator_through_report_plan_and_dashboard(call, tmp_path):
    call("project.create", {"path": str(tmp_path / "r.sme"), "name": "Formation", "model": {"id": "eu"}})
    root = call("framework.add", {"title": "Objectif global"})["created"]
    so = call("framework.add", {"parent_id": root, "title": "Objectif spécifique"})["created"]
    out = call("framework.add", {"parent_id": so, "title": "Les agents de santé sont formés"})["created"]
    spec = {
        "name": "Agents de santé formés",
        "unit": "count",
        "source_type": "routine",
        "aggregation": "sum",
        "categories": ["Femmes", "Hommes"],
        "target": 1500,
        "frequency": "quarterly",
    }
    ind = call("indicator.save", {"node_id": out, "indicator": spec})["saved"]
    values = [
        {"period": "2026-T1", "end": "2026-03-31", "by": {"Femmes": 180, "Hommes": 140}, "planned": 300},
        {"period": "2026-T2", "end": "2026-06-30", "by": {"Femmes": 300, "Hommes": 250}, "planned": 400},
    ]
    res = call("routine.save", {"indicator_id": ind, "values": values})
    an = res["analysis"]
    assert res["saved"] == 2 and an["current"] == 870 and an["status"] == "on_track" and res["project"]
    row = call("analysis.overview")["rows"][0]
    assert row["current"] == 870 and row["status"] == "on_track"
    assert not [i for i in call("report.checks")["items"] if i["severity"] == "warning"]

    report = tmp_path / "r.docx"
    call("report.generate", {"path": str(report), "lang": "fr"})
    text = "\n".join(p.text for p in Document(report).paragraphs)
    cells = [c.text for t in Document(report).tables for r in t.rows for c in r.cells]
    assert "870 en cumul au terme de la période 2026-T2" in text
    assert "Produit" in text  # the EU name of the level, not "Extrant"
    assert "2026-T1" in cells and "Prévu à date" in cells
    assert "Indicateur de suivi de routine" in text

    plan = call("report.plan", {"path": str(tmp_path / "p.xlsx"), "lang": "fr"})
    import polars as pl

    sheet = pl.read_excel(plan["path"], sheet_name="Plan de suivi", read_options={"header_row": 3})
    first = sheet.row(0, named=True)
    assert first["Niveau"] == "Produit" and first["Mode de calcul"].startswith("Suivi de routine")

    board = call("report.share", {"path": str(tmp_path / "b.html"), "lang": "fr"})
    html = open(board["path"], encoding="utf-8").read()
    assert "2026-T2" in html and "Produit" in html
